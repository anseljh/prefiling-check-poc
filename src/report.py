import docx
from typing import Dict, Any

def generate_report(results: Dict[str, Any], output_path: str):
    doc = docx.Document()
    doc.add_heading("Citation Check Report", 0)
    
    total = results.get("total", 0)
    verified = len(results.get("verified_correct", []))
    incorrect = len(results.get("verified_incorrect", []))
    unverified = len(results.get("not_verified", []))
    
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(f"Total Citations Found: {total}")
    doc.add_paragraph(f"Verified Correct: {verified}")
    doc.add_paragraph(f"Verified Incorrect: {incorrect}")
    doc.add_paragraph(f"Not Verified: {unverified}")
    
    if results.get("verified_incorrect"):
        doc.add_heading("Verified Incorrect Citations", level=1)
        for item in results["verified_incorrect"]:
            p = doc.add_paragraph()
            p.add_run(f"Citation: {item['citation']}").bold = True
            doc.add_paragraph(f"Context: {item['context']}")
            doc.add_paragraph(f"Expected Case: {item['expected']}")
            for reason in item['reasons']:
                doc.add_paragraph(f"- {reason}", style='List Bullet')
                
    if results.get("not_verified"):
        doc.add_heading("Unverified Citations", level=1)
        for item in results["not_verified"]:
            p = doc.add_paragraph()
            p.add_run(f"Citation: {item['citation']}").bold = True
            doc.add_paragraph(f"Context: {item['context']}")
            doc.add_paragraph(f"Reason: {item['reason']}")
            
    if results.get("verified_correct"):
        doc.add_heading("Verified Correct Citations", level=1)
        for item in results["verified_correct"]:
            p = doc.add_paragraph()
            p.add_run(f"Citation: {item['citation']}").bold = True
            doc.add_paragraph(f"Case: {item['case_name']}")
            
    doc.save(output_path)
